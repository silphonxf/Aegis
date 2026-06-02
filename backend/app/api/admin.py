from typing import Any, Dict, List, Optional, Set, Tuple
from datetime import datetime
from io import BytesIO, StringIO
import csv

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.api.deps import require_roles
from app.core.logging import get_logger
from app.core.security import get_password_hash
from app.db.session import get_db
from app.models.asset import Asset
from app.models.audit import AuditLog
from app.models.inspection import InspectionPoint
from app.models.shared_data import Room
from app.models.system import System, SystemUserBinding
from app.models.user import Role, User
from app.schemas.admin import (
    BatchCreateAssetsRequest,
    CreateAssetRequest,
    CreateInspectionPointRequest,
    CreateRoomRequest,
    CreateUserRequest,
    ThreatIntelBlockRequest,
    ThreatIntelQueryRequest,
    ThreatIntelQuickInputRequest,
)
from app.schemas.emergency_config import (
    EmergencyDbActionSaveRequest,
    EmergencyProcessActionSaveRequest,
    EmergencyServerActionSaveRequest,
    EmergencySshHostSaveRequest,
)
from app.services.emergency_config import (
    import_emergency_json_to_db,
    list_public_config,
    upsert_db_action,
    upsert_process_action,
    upsert_server_action,
    upsert_ssh_host,
)
from app.services.threatbook import ThreatbookError, batch_query_ip_reputation
from app.schemas.system import SystemCreate
from app.services.audit import log_action

router = APIRouter(prefix="/admin", tags=["admin"])
logger = get_logger("admin")


IP_HEADER_CANDIDATES = {"ip", "ip地址", "ip_address", "地址", "目标ip", "ipv4", "ipv6"}


def _extract_ips_from_excel(content: bytes) -> List[str]:
    try:
        workbook = load_workbook(filename=BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail={"code": "INVALID_EXCEL", "message": f"Excel 解析失败：{exc}"})

    ips: List[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        ip_col = None
        for idx, name in enumerate(header):
            if name in IP_HEADER_CANDIDATES:
                ip_col = idx
                break
        if ip_col is None:
            ip_col = 0
        for row in rows[1:]:
            if ip_col < len(row) and row[ip_col] is not None:
                ips.append(str(row[ip_col]).strip())
    return ips


@router.get("/users")
def list_users(
    page: int = 1,
    size: int = 20,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    logger.info("查询用户列表: page=%s size=%s", page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(User)
    total = q.count()
    users = q.offset((page - 1) * size).limit(size).all()
    logger.info("查询用户列表完成: total=%s returned=%s", total, len(users))
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {"id": u.id, "username": u.username, "role": u.role.code, "role_code": u.role.code, "is_active": u.is_active}
            for u in users
        ],
    }


@router.post("/users")
def create_user(
    payload: CreateUserRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    logger.info("创建用户请求: operator=%s username=%s role_code=%s", current_user.username, payload.username, payload.role_code)
    if db.query(User).filter(User.username == payload.username).first():
        raise HTTPException(status_code=400, detail={"code": "USER_EXISTS", "message": "用户名已存在"})
    role = db.query(Role).filter(Role.code == payload.role_code).first()
    if not role:
        logger.warning("创建用户失败: role_code=%s 不存在", payload.role_code)
        raise HTTPException(status_code=400, detail={"code": "ROLE_NOT_FOUND", "message": "角色不存在"})
    user = User(username=payload.username, password_hash=get_password_hash(payload.password), role_id=role.id)
    db.add(user)
    db.commit()
    db.refresh(user)
    logger.info("创建用户成功: user_id=%s username=%s", user.id, user.username)
    log_action(db, "create_user", "user", current_user, {"new_user_id": user.id, "username": user.username})
    return {"id": user.id}


@router.get("/systems")
def list_systems(
    page: int = 1,
    size: int = 20,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("查询系统列表: page=%s size=%s", page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(System)
    total = q.count()
    items = q.offset((page - 1) * size).limit(size).all()
    logger.info("查询系统列表完成: total=%s returned=%s", total, len(items))
    binding_rows = db.query(SystemUserBinding).filter(SystemUserBinding.binding_role == "owner").all()
    owner_map = {}
    for row in binding_rows:
        owner_map.setdefault(row.system_id, []).append(row.user_id)
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {
                "id": i.id,
                "system_code": i.system_code,
                "name": i.name,
                "env": i.env,
                "owner_user_id": i.owner_user_id,
                "owner_user_ids": owner_map.get(i.id, []),
                "check_frequency": getattr(i, "check_frequency", None),
                "remark": getattr(i, "remark", None),
            }
            for i in items
        ],
    }


@router.post("/systems")
def create_system(
    payload: SystemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("创建系统请求: operator=%s system_code=%s", current_user.username, payload.system_code)
    if db.query(System).filter(System.system_code == payload.system_code).first():
        raise HTTPException(status_code=409, detail={"code": "SYSTEM_CODE_EXISTS", "message": "系统编号已存在"})

    primary_owner_id = payload.owner_user_id or (payload.owner_user_ids[0] if payload.owner_user_ids else None)
    system = System(
        system_code=payload.system_code,
        name=payload.name,
        owner_user_id=primary_owner_id,
        env=payload.env,
        check_frequency=payload.check_frequency,
        remark=payload.remark,
    )
    db.add(system)
    db.flush()

    owner_ids = []
    for uid in ([primary_owner_id] if primary_owner_id else []) + list(payload.owner_user_ids):
        if uid and uid not in owner_ids:
            owner_ids.append(uid)
    for index, uid in enumerate(owner_ids):
        db.add(SystemUserBinding(
            system_id=system.id,
            user_id=uid,
            binding_role="owner",
            is_primary=(index == 0),
        ))

    db.commit()
    db.refresh(system)
    logger.info("创建系统成功: system_id=%s system_code=%s", system.id, system.system_code)
    log_action(db, "create_system", "system", current_user, {"system_id": system.id, "system_code": system.system_code})
    return {"id": system.id}


@router.get("/audit-logs")
def list_audit_logs(
    page: int = 1,
    size: int = 20,
    action: Optional[str] = None,
    username: Optional[str] = None,
    resource: Optional[str] = None,
    start_at: Optional[datetime] = None,
    end_at: Optional[datetime] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    logger.info("查询审计日志: action=%s username=%s resource=%s page=%s size=%s", action, username, resource, page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(AuditLog)

    if action:
        q = q.filter(AuditLog.action == action)
    if username:
        q = q.filter(AuditLog.username == username)
    if resource:
        q = q.filter(AuditLog.resource == resource)
    if start_at:
        q = q.filter(AuditLog.created_at >= start_at)
    if end_at:
        q = q.filter(AuditLog.created_at <= end_at)
    if keyword:
        q = q.filter(AuditLog.detail.like(f"%{keyword}%"))

    total = q.count()
    items = q.order_by(AuditLog.created_at.desc()).offset((page - 1) * size).limit(size).all()
    logger.info("查询审计日志完成: total=%s returned=%s", total, len(items))
    return {
        "page": page,
        "size": size,
        "total": total,
        "filters": {
            "action": action,
            "username": username,
            "resource": resource,
            "start_at": start_at,
            "end_at": end_at,
            "keyword": keyword,
        },
        "items": [
            {
                "id": i.id,
                "user_id": i.user_id,
                "username": i.username,
                "action": i.action,
                "resource": i.resource,
                "detail": i.detail,
                "created_at": i.created_at,
            }
            for i in items
        ],
    }


def _asset_query(db: Session, system_id: Optional[int], category: Optional[str], status: Optional[str], keyword: Optional[str]):
    q = db.query(Asset)
    if system_id is not None:
        q = q.filter(Asset.system_id == system_id)
    if category:
        q = q.filter(Asset.category == category)
    if status:
        q = q.filter(Asset.status == status)
    if keyword:
        q = q.filter((Asset.asset_code.like(f"%{keyword}%")) | (Asset.name.like(f"%{keyword}%")))
    return q


@router.get("/assets")
def list_assets(
    page: int = 1,
    size: int = 20,
    system_id: Optional[int] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("查询资产列表: system_id=%s category=%s status=%s keyword=%s page=%s size=%s", system_id, category, status, keyword, page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)

    q = _asset_query(db, system_id, category, status, keyword)
    total = q.count()
    items = q.order_by(Asset.id.desc()).offset((page - 1) * size).limit(size).all()
    logger.info("查询资产列表完成: total=%s returned=%s", total, len(items))
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {
                "id": a.id,
                "asset_code": a.asset_code,
                "name": a.name,
                "category": a.category,
                "system_id": a.system_id,
                "room_id": a.room_id,
                "location": a.location,
                "ip_address": a.ip_address,
                "port": a.port,
                "connection_type": a.connection_type,
                "status": a.status,
                "remark": a.remark,
                "created_at": a.created_at,
                "updated_at": a.updated_at,
            }
            for a in items
        ],
    }


@router.get("/assets/export")
def export_assets_csv(
    system_id: Optional[int] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    q = _asset_query(db, system_id, category, status, keyword)
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "id",
        "asset_code",
        "name",
        "category",
        "system_id",
        "room_id",
        "location",
        "ip_address",
        "port",
        "connection_type",
        "status",
        "remark",
        "created_at",
    ])
    for a in q.order_by(Asset.id.desc()).all():
        writer.writerow([
            a.id,
            a.asset_code,
            a.name,
            a.category,
            a.system_id or "",
            a.room_id or "",
            a.location or "",
            a.ip_address or "",
            a.port or "",
            a.connection_type or "",
            a.status,
            a.remark or "",
            a.created_at,
        ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=assets.csv"},
    )


@router.get("/assets/summary")
def assets_summary(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    total = db.query(func.count(Asset.id)).scalar() or 0
    by_status_rows = db.query(Asset.status, func.count(Asset.id)).group_by(Asset.status).all()
    by_category_rows = db.query(Asset.category, func.count(Asset.id)).group_by(Asset.category).all()

    return {
        "total": total,
        "by_status": [{"status": row[0], "count": row[1]} for row in by_status_rows],
        "by_category": [{"category": row[0], "count": row[1]} for row in by_category_rows],
    }


@router.post("/assets")
def create_asset(
    payload: CreateAssetRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("创建资产请求: operator=%s asset_code=%s", current_user.username, payload.asset_code)
    if db.query(Asset).filter(Asset.asset_code == payload.asset_code).first():
        raise HTTPException(status_code=400, detail={"code": "ASSET_EXISTS", "message": "资产编码已存在"})

    asset = Asset(**payload.dict())
    db.add(asset)
    db.commit()
    db.refresh(asset)
    logger.info("创建资产成功: asset_id=%s asset_code=%s", asset.id, asset.asset_code)
    log_action(db, "create_asset", "asset", current_user, {"asset_id": asset.id, "asset_code": asset.asset_code})
    return {"id": asset.id}


@router.get("/rooms")
def list_rooms(
    page: int = 1,
    size: int = 200,
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    page = max(page, 1)
    size = min(max(size, 1), 500)
    q = db.query(Room)
    if not include_inactive:
        q = q.filter(Room.is_active.is_(True))
    total = q.count()
    items = q.order_by(Room.id.asc()).offset((page - 1) * size).limit(size).all()
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {
                "id": item.id,
                "room_code": item.room_code,
                "room_name": item.room_name,
                "building": item.building,
                "floor": item.floor,
                "location_detail": item.location_detail,
                "remark": item.remark,
                "is_active": item.is_active,
                "created_at": item.created_at,
                "updated_at": item.updated_at,
            }
            for item in items
        ],
    }


@router.post("/rooms")
def create_room(
    payload: CreateRoomRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    if db.query(Room).filter(Room.room_code == payload.room_code).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_EXISTS", "message": "机房编码已存在"})
    room = Room(**payload.dict())
    db.add(room)
    db.commit()
    db.refresh(room)
    log_action(db, "create_room", "room", current_user, {"room_id": room.id, "room_code": room.room_code})
    return {"id": room.id}


@router.get("/inspection-points")
def list_inspection_points(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    items = db.query(InspectionPoint).order_by(InspectionPoint.id.asc()).all()
    return {
        "items": [
            {
                "id": item.id,
                "room_id": item.room_id,
                "system_id": item.system_id,
                "point_code": item.point_code,
                "point_name": item.point_name,
                "point_type": item.point_type,
                "qr_content": item.qr_content,
                "nfc_tag": item.nfc_tag,
                "location_detail": item.location_detail,
                "is_active": item.is_active,
            }
            for item in items
        ]
    }


@router.post("/inspection-points")
def create_inspection_point(
    payload: CreateInspectionPointRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    if db.query(InspectionPoint).filter(InspectionPoint.point_code == payload.point_code).first():
        raise HTTPException(status_code=400, detail={"code": "POINT_EXISTS", "message": "巡检点编码已存在"})
    if not db.query(Room).filter(Room.id == payload.room_id, Room.is_active.is_(True)).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在或已停用"})
    if payload.system_id is not None and not db.query(System).filter(System.id == payload.system_id).first():
        raise HTTPException(status_code=400, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})
    point = InspectionPoint(
        room_id=payload.room_id,
        system_id=payload.system_id,
        point_code=payload.point_code,
        point_name=payload.point_name,
        point_type=payload.point_type,
        qr_content=payload.qr_content or payload.point_code,
        nfc_tag=payload.nfc_tag,
        location_detail=payload.location_detail,
        is_active=payload.is_active,
        location=payload.location_detail,
    )
    db.add(point)
    db.commit()
    db.refresh(point)
    log_action(db, "create_inspection_point", "inspection_point", current_user, {"point_id": point.id, "point_code": point.point_code})
    return {"id": point.id}


@router.post("/assets/batch")
def batch_create_assets(
    payload: BatchCreateAssetsRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("批量创建资产请求: operator=%s item_count=%s", current_user.username, len(payload.items))
    created = []
    skipped = []

    seen_codes: Set[str] = set()
    items_to_create = []
    for item in payload.items:
        code = item.asset_code.strip()
        if code in seen_codes:
            skipped.append({"asset_code": code, "reason": "BATCH_DUPLICATE"})
            continue
        seen_codes.add(code)
        items_to_create.append(item)

    existing_codes = set()
    if items_to_create:
        existing_codes = {
            row[0]
            for row in db.query(Asset.asset_code).filter(Asset.asset_code.in_([it.asset_code for it in items_to_create])).all()
        }

    for item in items_to_create:
        if item.asset_code in existing_codes:
            skipped.append({"asset_code": item.asset_code, "reason": "ASSET_EXISTS"})
            continue
        asset = Asset(**item.dict())
        db.add(asset)
        db.flush()
        created.append({"id": asset.id, "asset_code": asset.asset_code})

    db.commit()
    logger.info("批量创建资产完成: created=%s skipped=%s", len(created), len(skipped))
    log_action(
        db,
        "batch_create_assets",
        "asset",
        current_user,
        {"created_count": len(created), "skipped_count": len(skipped)},
    )
    return {"created": created, "skipped": skipped}


@router.post("/threat-intel/ip-reputation")
def query_ip_reputation(
    payload: ThreatIntelQueryRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("威胁情报批量查询请求: operator=%s count=%s", current_user.username, len(payload.ips))
    try:
        result = batch_query_ip_reputation(payload.ips, lang=payload.lang, realtime_verdict=payload.realtime_verdict)
    except ThreatbookError as exc:
        logger.warning("威胁情报批量查询失败: %s", str(exc))
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    logger.info("威胁情报批量查询完成: total=%s high_risk=%s", result["summary"]["total"], result["summary"]["high_risk"])
    log_action(
        db,
        "query_ip_reputation",
        "threat_intel",
        current_user,
        {
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.post("/threat-intel/ip-reputation/quick")
def query_ip_reputation_quick(
    payload: ThreatIntelQuickInputRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    raw_items = [payload.raw_input]
    try:
        result = batch_query_ip_reputation(raw_items, lang=payload.lang, realtime_verdict=payload.realtime_verdict)
    except ThreatbookError as exc:
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    log_action(
        db,
        "query_ip_reputation_quick",
        "threat_intel",
        current_user,
        {
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.post("/threat-intel/ip-reputation/excel")
async def query_ip_reputation_excel(
    file: UploadFile = File(...),
    lang: str = "zh",
    realtime_verdict: bool = True,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(status_code=400, detail={"code": "INVALID_FILE_TYPE", "message": "仅支持上传 .xlsx 类 Excel 文件"})

    content = await file.read()
    raw_ips = _extract_ips_from_excel(content)
    try:
        result = batch_query_ip_reputation(raw_ips, lang=lang, realtime_verdict=realtime_verdict)
    except ThreatbookError as exc:
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    result["import"] = {"filename": filename, "source": "excel"}
    log_action(
        db,
        "query_ip_reputation_excel",
        "threat_intel",
        current_user,
        {
            "filename": filename,
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.get("/emergency-config")
def get_emergency_config(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    return list_public_config(db)


@router.post("/emergency-config/ssh-hosts")
def save_emergency_ssh_host(
    payload: EmergencySshHostSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_ssh_host(payload, db=db)
    log_action(db, "save_emergency_ssh_host", "emergency_config", current_user, {"host_code": payload.host_code})
    return item


@router.post("/emergency-config/server-actions")
def save_emergency_server_action(
    payload: EmergencyServerActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_server_action(payload, db=db)
    log_action(db, "save_emergency_server_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.post("/emergency-config/database-actions")
def save_emergency_db_action(
    payload: EmergencyDbActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_db_action(payload, db=db)
    log_action(db, "save_emergency_db_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.post("/emergency-config/process-actions")
def save_emergency_process_action(
    payload: EmergencyProcessActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_process_action(payload, db=db)
    log_action(db, "save_emergency_process_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.post("/emergency-config/import-json")
def import_emergency_config_json(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = import_emergency_json_to_db(db)
    log_action(db, "import_emergency_config_json", "emergency_config", current_user, result)
    return result


@router.post("/threat-intel/block-ip")
def block_high_risk_ip(
    payload: ThreatIntelBlockRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    result = {
        "status": "mocked",
        "dry_run": payload.dry_run,
        "ip": payload.ip,
        "risk_level": payload.risk_level,
        "reason": payload.reason,
        "source": payload.source,
        "message": "已预留防火墙封禁接口，当前为 mock 返回，后续可替换为真实执行器。",
    }
    log_action(db, "block_ip_request", "threat_intel", current_user, result)
    return result
