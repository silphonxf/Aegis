from typing import Any, Dict, List, Optional, Set, Tuple
from datetime import datetime
from io import BytesIO, StringIO
import csv
import hashlib
import json
import secrets

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.api.deps import require_roles
from app.core.logging import get_logger
from app.core.security import get_password_hash
from app.db.session import get_db
from app.models.asset import Asset
from app.models.audit import AuditLog
from app.models.ai_external_key import AIExternalApiKey
from app.models.inspection import InspectionPoint, InspectionRecord
from app.models.shared_data import Room
from app.models.system import System, SystemLogConfig, SystemUserBinding
from app.models.user import Role, User
from app.schemas.admin import (
    AIEngineChatTestRequest,
    AIEngineConfigRequest,
    BatchCreateAssetsRequest,
    CreateAIExternalApiKeyRequest,
    CreateAssetRequest,
    CreateInspectionPointRequest,
    CreateRoomRequest,
    CreateUserRequest,
    FirewallBlockConfigRequest,
    ThreatIntelBlockRequest,
    ThreatIntelQueryRequest,
    ThreatIntelQuickInputRequest,
    UpdateAssetRequest,
    UpdateInspectionPointRequest,
    UpdateRoomRequest,
)
from app.schemas.ai import ChatRequest
from app.services.ai_engine_config import get_runtime_ai_config, serialize_ai_engine_config, upsert_ai_engine_config
from app.services.ai_provider import run_chat
from app.schemas.emergency_config import (
    EmergencyDbActionSaveRequest,
    EmergencyProcessActionSaveRequest,
    EmergencyServerActionSaveRequest,
    EmergencySshHostSaveRequest,
)
from app.services.emergency_config import (
    delete_db_action,
    delete_process_action,
    delete_server_action,
    delete_ssh_host,
    import_emergency_json_to_db,
    list_public_config,
    upsert_db_action,
    upsert_process_action,
    upsert_server_action,
    upsert_ssh_host,
)
from app.services.threatbook import ThreatbookError, batch_query_ip_reputation
from app.services.firewall import FirewallClientError, FirewallConfigError, FirewallError, block_ip_with_firewall
from app.services.firewall_config import serialize_firewall_config, upsert_firewall_config
from app.schemas.system import SystemCreate, SystemUpdate
from app.services.audit import log_action
from app.services.cache import delete_prefix, get_json, set_json

router = APIRouter(prefix="/admin", tags=["admin"])
logger = get_logger("admin")


IP_HEADER_CANDIDATES = {"ip", "ip地址", "ip_address", "地址", "目标ip", "ipv4", "ipv6"}
SYSTEM_LIST_CACHE_PREFIX = "aegis:cache:admin:systems:"


def _hash_external_api_key(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _serialize_ai_external_key(item: AIExternalApiKey) -> Dict[str, Any]:
    return {
        "id": item.id,
        "name": item.name,
        "key_prefix": item.key_prefix,
        "is_active": item.is_active,
        "remark": item.remark,
        "created_by": item.created_by,
        "created_at": item.created_at,
        "updated_at": item.updated_at,
        "last_used_at": item.last_used_at,
    }


def _extract_ips_from_excel(content: bytes) -> List[str]:
    try:
        workbook = load_workbook(filename=BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail={"code": "INVALID_EXCEL", "message": f"Excel 解析失败：{exc}"})

    ips: List[str] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        ip_col = None
        for idx, name in enumerate(header):
            if name in IP_HEADER_CANDIDATES:
                ip_col = idx
                break
        if ip_col is None:
            ip_col = 0
        for row in rows[1:]:
            if ip_col < len(row) and row[ip_col] is not None:
                ips.append(str(row[ip_col]).strip())
    return ips


@router.get("/users")
def list_users(
    page: int = 1,
    size: int = 20,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    logger.info("查询用户列表: page=%s size=%s", page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(User)
    total = q.count()
    users = q.offset((page - 1) * size).limit(size).all()
    logger.info("查询用户列表完成: total=%s returned=%s", total, len(users))
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {"id": u.id, "username": u.username, "role": u.role.code, "role_code": u.role.code, "is_active": u.is_active}
            for u in users
        ],
    }


@router.post("/users")
def create_user(
    payload: CreateUserRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    logger.info("创建用户请求: operator=%s username=%s role_code=%s", current_user.username, payload.username, payload.role_code)
    if db.query(User).filter(User.username == payload.username).first():
        raise HTTPException(status_code=400, detail={"code": "USER_EXISTS", "message": "用户名已存在"})
    role = db.query(Role).filter(Role.code == payload.role_code).first()
    if not role:
        logger.warning("创建用户失败: role_code=%s 不存在", payload.role_code)
        raise HTTPException(status_code=400, detail={"code": "ROLE_NOT_FOUND", "message": "角色不存在"})
    user = User(username=payload.username, password_hash=get_password_hash(payload.password), role_id=role.id)
    db.add(user)
    db.commit()
    db.refresh(user)
    logger.info("创建用户成功: user_id=%s username=%s", user.id, user.username)
    log_action(db, "create_user", "user", current_user, {"new_user_id": user.id, "username": user.username})
    return {"id": user.id}


@router.get("/ai-external-keys")
def list_ai_external_api_keys(
    page: int = 1,
    size: int = 50,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(AIExternalApiKey)
    total = q.count()
    rows = q.order_by(AIExternalApiKey.id.desc()).offset((page - 1) * size).limit(size).all()
    return {"page": page, "size": size, "total": total, "items": [_serialize_ai_external_key(item) for item in rows]}


@router.post("/ai-external-keys")
def create_ai_external_api_key(
    payload: CreateAIExternalApiKeyRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    api_key = f"aegis_ai_{secrets.token_urlsafe(32)}"
    row = AIExternalApiKey(
        name=payload.name.strip(),
        key_prefix=api_key[:18],
        key_hash=_hash_external_api_key(api_key),
        remark=payload.remark,
        created_by=current_user.username,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    log_action(db, "create_ai_external_api_key", "ai_external_key", current_user, {"key_id": row.id, "name": row.name, "key_prefix": row.key_prefix})
    data = _serialize_ai_external_key(row)
    data["apikey"] = api_key
    return data


@router.patch("/ai-external-keys/{key_id}/active")
def set_ai_external_api_key_active(
    key_id: int,
    is_active: bool,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    row = db.query(AIExternalApiKey).filter(AIExternalApiKey.id == key_id).first()
    if not row:
        raise HTTPException(status_code=404, detail={"code": "AI_EXTERNAL_KEY_NOT_FOUND", "message": "AI 对外接口 Key 不存在"})
    row.is_active = is_active
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    log_action(db, "set_ai_external_api_key_active", "ai_external_key", current_user, {"key_id": row.id, "is_active": is_active})
    return _serialize_ai_external_key(row)


@router.get("/ai-engine-config")
def get_ai_engine_config(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    return serialize_ai_engine_config(db)


@router.put("/ai-engine-config")
def save_ai_engine_config(
    payload: AIEngineConfigRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    result = upsert_ai_engine_config(db, payload)
    log_action(
        db,
        "save_ai_engine_config",
        "ai_engine_config",
        current_user,
        {
            "engine_type": result["engine_type"],
            "base_url": result["base_url"],
            "model": result["model"],
            "has_api_key": result["has_api_key"],
        },
    )
    return result


@router.post("/ai-engine-config/test-chat")
def test_ai_engine_chat(
    payload: AIEngineChatTestRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    runtime_config = payload.dict(
        exclude={"message", "conversation_id", "history"},
    )
    if payload.api_key is None:
        runtime_config["api_key"] = get_runtime_ai_config().get("api_key", "")

    result = run_chat(
        ChatRequest(message=payload.message, conversation_id=payload.conversation_id),
        history=payload.history[-12:],
        runtime_config=runtime_config,
    )
    result["tested_engine_type"] = payload.engine_type
    log_action(
        db,
        "test_ai_engine_chat",
        "ai_engine_config",
        current_user,
        {
            "engine_type": payload.engine_type,
            "base_url": payload.base_url,
            "model": payload.model,
            "mode": result.get("mode"),
            "elapsed_ms": result.get("elapsed_ms"),
            "fallback_reason": result.get("fallback_reason"),
        },
    )
    return result


@router.get("/systems")
def list_systems(
    page: int = 1,
    size: int = 20,
    keyword: Optional[str] = None,
    env: Optional[str] = None,
    owner_user_id: Optional[int] = None,
    is_active: Optional[bool] = None,
    sort_by: str = "id",
    sort_order: str = "asc",
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("查询系统列表: page=%s size=%s keyword=%s env=%s owner_user_id=%s", page, size, keyword, env, owner_user_id)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    cache_key = _system_list_cache_key(page, size, keyword, env, owner_user_id, is_active, sort_by, sort_order)
    cached = get_json(cache_key)
    if cached is not None:
        return cached
    q = db.query(System)
    if keyword:
        like = f"%{keyword}%"
        q = q.filter(or_(System.system_code.like(like), System.name.like(like), System.host_address.like(like)))
    if env:
        q = q.filter(System.env == env)
    if is_active is not None:
        q = q.filter(System.is_active.is_(is_active))
    if owner_user_id is not None:
        q = q.join(SystemUserBinding, SystemUserBinding.system_id == System.id).filter(
            SystemUserBinding.binding_role == "owner",
            SystemUserBinding.user_id == owner_user_id,
        )

    sort_column = {
        "id": System.id,
        "system_code": System.system_code,
        "name": System.name,
        "host_address": System.host_address,
        "env": System.env,
        "updated_at": System.updated_at,
    }.get(sort_by, System.id)
    q = q.order_by(sort_column.desc() if sort_order == "desc" else sort_column.asc())
    total = q.count()
    items = q.offset((page - 1) * size).limit(size).all()
    logger.info("查询系统列表完成: total=%s returned=%s", total, len(items))
    owner_map = _system_owner_map(db)
    log_map = _system_log_map(db, [item.id for item in items])
    payload = {
        "page": page,
        "size": size,
        "total": total,
        "items": [_serialize_system(item, owner_map, log_map) for item in items],
    }
    set_json(cache_key, payload)
    return payload


@router.post("/systems")
def create_system(
    payload: SystemCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("创建系统请求: operator=%s system_code=%s", current_user.username, payload.system_code)
    if db.query(System).filter(System.system_code == payload.system_code).first():
        raise HTTPException(status_code=409, detail={"code": "SYSTEM_CODE_EXISTS", "message": "系统编号已存在"})

    primary_owner_id = payload.owner_user_id or (payload.owner_user_ids[0] if payload.owner_user_ids else None)
    system = System(
        system_code=payload.system_code,
        name=payload.name,
        host_address=payload.host_address,
        owner_user_id=primary_owner_id,
        env=payload.env,
        check_frequency=payload.check_frequency,
        selfcheck_skill=payload.selfcheck_skill,
        remark=payload.remark,
    )
    db.add(system)
    db.flush()

    _replace_system_owner_bindings(db, system.id, _normalize_owner_ids(primary_owner_id, payload.owner_user_ids))
    _replace_system_log_configs(db, system.id, payload.log_configs)

    db.commit()
    db.refresh(system)
    logger.info("创建系统成功: system_id=%s system_code=%s", system.id, system.system_code)
    log_action(db, "create_system", "system", current_user, {"system_id": system.id, "system_code": system.system_code})
    _invalidate_system_reference_cache()
    return {"id": system.id}


@router.put("/systems/{system_id}")
def update_system(
    system_id: int,
    payload: SystemUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    system = db.query(System).filter(System.id == system_id).first()
    if not system:
        raise HTTPException(status_code=404, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})
    values = payload.dict(exclude_unset=True)
    if "system_code" in values and values["system_code"] != system.system_code:
        if db.query(System).filter(System.system_code == values["system_code"], System.id != system_id).first():
            raise HTTPException(status_code=409, detail={"code": "SYSTEM_CODE_EXISTS", "message": "系统编号已存在"})

    log_configs_value = values.pop("log_configs", None)
    owner_ids_value = values.pop("owner_user_ids", None)
    owner_user_id_value = values.get("owner_user_id") if owner_ids_value is not None else values.get("owner_user_id", system.owner_user_id)
    if owner_ids_value is not None or "owner_user_id" in values:
        owner_ids = _normalize_owner_ids(owner_user_id_value, owner_ids_value or [])
        values["owner_user_id"] = owner_ids[0] if owner_ids else None
        _replace_system_owner_bindings(db, system.id, owner_ids)
    if log_configs_value is not None:
        _replace_system_log_configs(db, system.id, log_configs_value)

    _apply_updates(system, values)
    db.commit()
    db.refresh(system)
    log_action(db, "update_system", "system", current_user, {"system_id": system.id, "system_code": system.system_code})
    _invalidate_system_reference_cache()
    return {"id": system.id}


@router.delete("/systems/{system_id}")
def deactivate_system(
    system_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    system = db.query(System).filter(System.id == system_id).first()
    if not system:
        raise HTTPException(status_code=404, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})
    system.is_active = False
    system.updated_at = datetime.utcnow()
    db.commit()
    log_action(db, "deactivate_system", "system", current_user, {"system_id": system.id, "system_code": system.system_code})
    _invalidate_system_reference_cache()
    return {"id": system.id, "is_active": system.is_active}


@router.patch("/systems/{system_id}/active")
def set_system_active(
    system_id: int,
    is_active: bool,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    system = db.query(System).filter(System.id == system_id).first()
    if not system:
        raise HTTPException(status_code=404, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})
    system.is_active = is_active
    system.updated_at = datetime.utcnow()
    db.commit()
    log_action(db, "set_system_active", "system", current_user, {"system_id": system.id, "system_code": system.system_code, "is_active": is_active})
    _invalidate_system_reference_cache()
    return {"id": system.id, "is_active": system.is_active}


@router.get("/audit-logs")
def list_audit_logs(
    page: int = 1,
    size: int = 20,
    action: Optional[str] = None,
    username: Optional[str] = None,
    resource: Optional[str] = None,
    start_at: Optional[datetime] = None,
    end_at: Optional[datetime] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    logger.info("查询审计日志: action=%s username=%s resource=%s page=%s size=%s", action, username, resource, page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(AuditLog)

    if action:
        q = q.filter(AuditLog.action == action)
    if username:
        q = q.filter(AuditLog.username == username)
    if resource:
        q = q.filter(AuditLog.resource == resource)
    if start_at:
        q = q.filter(AuditLog.created_at >= start_at)
    if end_at:
        q = q.filter(AuditLog.created_at <= end_at)
    if keyword:
        q = q.filter(AuditLog.detail.like(f"%{keyword}%"))

    total = q.count()
    items = q.order_by(AuditLog.created_at.desc()).offset((page - 1) * size).limit(size).all()
    logger.info("查询审计日志完成: total=%s returned=%s", total, len(items))
    return {
        "page": page,
        "size": size,
        "total": total,
        "filters": {
            "action": action,
            "username": username,
            "resource": resource,
            "start_at": start_at,
            "end_at": end_at,
            "keyword": keyword,
        },
        "items": [
            {
                "id": i.id,
                "user_id": i.user_id,
                "username": i.username,
                "action": i.action,
                "resource": i.resource,
                "detail": i.detail,
                "created_at": i.created_at,
            }
            for i in items
        ],
    }


def _asset_query(db: Session, system_id: Optional[int], category: Optional[str], status: Optional[str], keyword: Optional[str]):
    q = db.query(Asset)
    if system_id is not None:
        q = q.filter(Asset.system_id == system_id)
    if category:
        q = q.filter(Asset.category == category)
    if status:
        q = q.filter(Asset.status == status)
    if keyword:
        q = q.filter((Asset.asset_code.like(f"%{keyword}%")) | (Asset.name.like(f"%{keyword}%")))
    return q


def _validate_asset_refs(db: Session, system_id: Optional[int], room_id: Optional[int]) -> None:
    if system_id is not None and not db.query(System).filter(System.id == system_id).first():
        raise HTTPException(status_code=400, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})
    if room_id is not None and not db.query(Room).filter(Room.id == room_id, Room.is_active.is_(True)).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在或已停用"})


def _validate_inspection_point_refs(db: Session, room_id: Optional[int], system_id: Optional[int]) -> None:
    if room_id is not None and not db.query(Room).filter(Room.id == room_id, Room.is_active.is_(True)).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在或已停用"})
    if system_id is not None and not db.query(System).filter(System.id == system_id).first():
        raise HTTPException(status_code=400, detail={"code": "SYSTEM_NOT_FOUND", "message": "系统不存在"})


def _normalize_check_items(items: Optional[List[str]]) -> str:
    cleaned = [str(item).strip() for item in (items or []) if str(item).strip()]
    return json.dumps(cleaned, ensure_ascii=False)


def _parse_check_items(value: Optional[str]) -> List[str]:
    if not value:
        return []
    try:
        data = json.loads(value)
        if isinstance(data, list):
            return [str(item).strip() for item in data if str(item).strip()]
    except Exception:
        pass
    return [line.strip() for line in value.splitlines() if line.strip()]


def _generate_room_code(db: Session) -> str:
    prefix = datetime.utcnow().strftime("ROOM-%Y%m%d%H%M%S")
    candidate = prefix
    index = 1
    while db.query(Room).filter(Room.room_code == candidate).first():
        index += 1
        candidate = f"{prefix}-{index}"
    return candidate


def _ensure_room_inspection_point(db: Session, room: Room) -> None:
    qr_content = (getattr(room, "qr_content", None) or "").strip()
    if not qr_content:
        return
    point = (
        db.query(InspectionPoint)
        .filter(InspectionPoint.room_id == room.id, InspectionPoint.point_code == f"ROOM-{room.id}")
        .first()
    )
    if not point:
        point = db.query(InspectionPoint).filter(InspectionPoint.qr_content == qr_content).first()
    if not point:
        point = InspectionPoint(room_id=room.id, point_code=f"ROOM-{room.id}", is_active=True)
        db.add(point)
    point.room_id = room.id
    point.point_code = point.point_code or f"ROOM-{room.id}"
    point.point_name = room.room_name
    point.point_type = "room"
    point.qr_content = qr_content
    point.nfc_tag = (getattr(room, "nfc_tag", None) or None)
    point.location = room.room_name
    point.location_detail = room.location_detail or room.room_name
    point.is_active = room.is_active
    point.updated_at = datetime.utcnow()


def _apply_updates(instance, values: Dict[str, Any]) -> None:
    for key, value in values.items():
        setattr(instance, key, value)
    if hasattr(instance, "updated_at"):
        instance.updated_at = datetime.utcnow()


def _system_list_cache_key(
    page: int,
    size: int,
    keyword: Optional[str],
    env: Optional[str],
    owner_user_id: Optional[int],
    is_active: Optional[bool],
    sort_by: str,
    sort_order: str,
) -> str:
    payload = {
        "page": page,
        "size": size,
        "keyword": keyword or "",
        "env": env or "",
        "owner_user_id": owner_user_id,
        "is_active": is_active,
        "sort_by": sort_by,
        "sort_order": sort_order,
    }
    digest = hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()[:16]
    return f"{SYSTEM_LIST_CACHE_PREFIX}{digest}"


def _invalidate_system_reference_cache() -> None:
    delete_prefix(SYSTEM_LIST_CACHE_PREFIX)
    delete_prefix("aegis:cache:systems:")


def _normalize_owner_ids(owner_user_id: Optional[int], owner_user_ids: Optional[List[int]]) -> List[int]:
    owner_ids = []
    for uid in ([owner_user_id] if owner_user_id else []) + list(owner_user_ids or []):
        if uid and uid not in owner_ids:
            owner_ids.append(uid)
    return owner_ids


def _replace_system_owner_bindings(db: Session, system_id: int, owner_ids: List[int]) -> None:
    db.query(SystemUserBinding).filter(
        SystemUserBinding.system_id == system_id,
        SystemUserBinding.binding_role == "owner",
    ).delete()
    for index, uid in enumerate(owner_ids):
        db.add(SystemUserBinding(
            system_id=system_id,
            user_id=uid,
            binding_role="owner",
            is_primary=(index == 0),
        ))


def _replace_system_log_configs(db: Session, system_id: int, log_configs: List[Any]) -> None:
    db.query(SystemLogConfig).filter(SystemLogConfig.system_id == system_id).delete()
    for item in log_configs or []:
        getter = item.get if isinstance(item, dict) else lambda key, default=None: getattr(item, key, default)
        db.add(SystemLogConfig(
            system_id=system_id,
            log_name=getter("log_name"),
            absolute_path=getter("absolute_path"),
            log_level=getter("log_level", "warning"),
            is_active=getter("is_active", True),
            remark=getter("remark"),
        ))


def _system_owner_map(db: Session) -> Dict[int, List[int]]:
    binding_rows = db.query(SystemUserBinding).filter(SystemUserBinding.binding_role == "owner").all()
    owner_map: Dict[int, List[int]] = {}
    for row in binding_rows:
        owner_map.setdefault(row.system_id, []).append(row.user_id)
    return owner_map


def _system_log_map(db: Session, system_ids: List[int]) -> Dict[int, List[SystemLogConfig]]:
    if not system_ids:
        return {}
    rows = (
        db.query(SystemLogConfig)
        .filter(SystemLogConfig.system_id.in_(system_ids))
        .order_by(SystemLogConfig.id.asc())
        .all()
    )
    result: Dict[int, List[SystemLogConfig]] = {}
    for row in rows:
        result.setdefault(row.system_id, []).append(row)
    return result


def _serialize_log_config(item: SystemLogConfig) -> Dict[str, Any]:
    return {
        "id": item.id,
        "system_id": item.system_id,
        "log_name": item.log_name,
        "absolute_path": item.absolute_path,
        "log_level": item.log_level,
        "is_active": item.is_active,
        "remark": item.remark,
    }


def _serialize_system(item: System, owner_map: Dict[int, List[int]], log_map: Dict[int, List[SystemLogConfig]]) -> Dict[str, Any]:
    return {
        "id": item.id,
        "system_code": item.system_code,
        "name": item.name,
        "host_address": item.host_address,
        "env": item.env,
        "owner_user_id": item.owner_user_id,
        "owner_user_ids": owner_map.get(item.id, []),
        "check_frequency": getattr(item, "check_frequency", None),
        "selfcheck_skill": getattr(item, "selfcheck_skill", None),
        "remark": getattr(item, "remark", None),
        "is_active": getattr(item, "is_active", True),
        "log_configs": [_serialize_log_config(log) for log in log_map.get(item.id, [])],
    }


def _parse_inspection_note(value: Optional[str]) -> Dict[str, Any]:
    if not value:
        return {"note": "", "check_results": [], "monitoring_confirmation": None}
    try:
        data = json.loads(value)
        if isinstance(data, dict):
            return {
                "note": data.get("note") or "",
                "check_results": data.get("check_results") if isinstance(data.get("check_results"), list) else [],
                "monitoring_confirmation": data.get("monitoring_confirmation"),
            }
    except Exception:
        pass
    return {"note": value, "check_results": [], "monitoring_confirmation": None}


@router.get("/assets")
def list_assets(
    page: int = 1,
    size: int = 20,
    system_id: Optional[int] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("查询资产列表: system_id=%s category=%s status=%s keyword=%s page=%s size=%s", system_id, category, status, keyword, page, size)
    page = max(page, 1)
    size = min(max(size, 1), 100)

    q = _asset_query(db, system_id, category, status, keyword)
    total = q.count()
    items = q.order_by(Asset.id.desc()).offset((page - 1) * size).limit(size).all()
    logger.info("查询资产列表完成: total=%s returned=%s", total, len(items))
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {
                "id": a.id,
                "asset_code": a.asset_code,
                "name": a.name,
                "category": a.category,
                "system_id": a.system_id,
                "room_id": a.room_id,
                "location": a.location,
                "ip_address": a.ip_address,
                "port": a.port,
                "connection_type": a.connection_type,
                "status": a.status,
                "remark": a.remark,
                "created_at": a.created_at,
                "updated_at": a.updated_at,
            }
            for a in items
        ],
    }


@router.get("/assets/export")
def export_assets_csv(
    system_id: Optional[int] = None,
    category: Optional[str] = None,
    status: Optional[str] = None,
    keyword: Optional[str] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    q = _asset_query(db, system_id, category, status, keyword)
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "id",
        "asset_code",
        "name",
        "category",
        "system_id",
        "room_id",
        "location",
        "ip_address",
        "port",
        "connection_type",
        "status",
        "remark",
        "created_at",
    ])
    for a in q.order_by(Asset.id.desc()).all():
        writer.writerow([
            a.id,
            a.asset_code,
            a.name,
            a.category,
            a.system_id or "",
            a.room_id or "",
            a.location or "",
            a.ip_address or "",
            a.port or "",
            a.connection_type or "",
            a.status,
            a.remark or "",
            a.created_at,
        ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=assets.csv"},
    )


@router.get("/assets/summary")
def assets_summary(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    total = db.query(func.count(Asset.id)).scalar() or 0
    by_status_rows = db.query(Asset.status, func.count(Asset.id)).group_by(Asset.status).all()
    by_category_rows = db.query(Asset.category, func.count(Asset.id)).group_by(Asset.category).all()

    return {
        "total": total,
        "by_status": [{"status": row[0], "count": row[1]} for row in by_status_rows],
        "by_category": [{"category": row[0], "count": row[1]} for row in by_category_rows],
    }


@router.post("/assets")
def create_asset(
    payload: CreateAssetRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("创建资产请求: operator=%s asset_code=%s", current_user.username, payload.asset_code)
    if db.query(Asset).filter(Asset.asset_code == payload.asset_code).first():
        raise HTTPException(status_code=400, detail={"code": "ASSET_EXISTS", "message": "资产编码已存在"})
    _validate_asset_refs(db, payload.system_id, payload.room_id)

    asset = Asset(**payload.dict())
    db.add(asset)
    db.commit()
    db.refresh(asset)
    logger.info("创建资产成功: asset_id=%s asset_code=%s", asset.id, asset.asset_code)
    log_action(db, "create_asset", "asset", current_user, {"asset_id": asset.id, "asset_code": asset.asset_code})
    return {"id": asset.id}


@router.put("/assets/{asset_id}")
def update_asset(
    asset_id: int,
    payload: UpdateAssetRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail={"code": "ASSET_NOT_FOUND", "message": "资产不存在"})

    values = payload.dict(exclude_unset=True)
    if "asset_code" in values and values["asset_code"] != asset.asset_code:
        if db.query(Asset).filter(Asset.asset_code == values["asset_code"], Asset.id != asset_id).first():
            raise HTTPException(status_code=400, detail={"code": "ASSET_EXISTS", "message": "资产编码已存在"})
    if "system_id" in values or "room_id" in values:
        _validate_asset_refs(db, values.get("system_id", asset.system_id), values.get("room_id", asset.room_id))

    _apply_updates(asset, values)
    db.commit()
    db.refresh(asset)
    log_action(db, "update_asset", "asset", current_user, {"asset_id": asset.id, "asset_code": asset.asset_code})
    return {"id": asset.id}


@router.delete("/assets/{asset_id}")
def retire_asset(
    asset_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail={"code": "ASSET_NOT_FOUND", "message": "资产不存在"})
    asset.status = "retired"
    asset.updated_at = datetime.utcnow()
    db.commit()
    log_action(db, "retire_asset", "asset", current_user, {"asset_id": asset.id, "asset_code": asset.asset_code})
    return {"id": asset.id, "status": asset.status}


@router.get("/rooms")
def list_rooms(
    page: int = 1,
    size: int = 200,
    keyword: Optional[str] = None,
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    page = max(page, 1)
    size = min(max(size, 1), 500)
    q = db.query(Room)
    if not include_inactive:
        q = q.filter(Room.is_active.is_(True))
    if keyword:
        q = q.filter(
            (Room.room_code.like(f"%{keyword}%"))
            | (Room.room_name.like(f"%{keyword}%"))
            | (Room.qr_content.like(f"%{keyword}%"))
            | (Room.nfc_tag.like(f"%{keyword}%"))
        )
    total = q.count()
    items = q.order_by(Room.id.asc()).offset((page - 1) * size).limit(size).all()
    return {
        "page": page,
        "size": size,
        "total": total,
        "items": [
            {
                "id": item.id,
                "room_code": item.room_code,
                "room_name": item.room_name,
                "qr_content": getattr(item, "qr_content", None),
                "nfc_tag": getattr(item, "nfc_tag", None),
                "check_items": _parse_check_items(getattr(item, "check_items", None)),
                "weekday_inspection_count": getattr(item, "weekday_inspection_count", 1),
                "holiday_inspection_count": getattr(item, "holiday_inspection_count", 1),
                "building": item.building,
                "floor": item.floor,
                "location_detail": item.location_detail,
                "remark": item.remark,
                "is_active": item.is_active,
                "created_at": item.created_at,
                "updated_at": item.updated_at,
            }
            for item in items
        ],
    }


@router.post("/rooms")
def create_room(
    payload: CreateRoomRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    room_code = payload.room_code or _generate_room_code(db)
    if db.query(Room).filter(Room.room_code == room_code).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_EXISTS", "message": "机房编码已存在"})
    if db.query(Room).filter(Room.qr_content == payload.qr_content).first() or db.query(InspectionPoint).filter(InspectionPoint.qr_content == payload.qr_content).first():
        raise HTTPException(status_code=400, detail={"code": "ROOM_QR_EXISTS", "message": "二维码值已存在"})
    if payload.nfc_tag and (db.query(Room).filter(Room.nfc_tag == payload.nfc_tag).first() or db.query(InspectionPoint).filter(InspectionPoint.nfc_tag == payload.nfc_tag).first()):
        raise HTTPException(status_code=400, detail={"code": "ROOM_NFC_EXISTS", "message": "NFC 值已存在"})
    values = payload.dict()
    values["room_code"] = room_code
    values["check_items"] = _normalize_check_items(payload.check_items)
    room = Room(**values)
    db.add(room)
    db.flush()
    _ensure_room_inspection_point(db, room)
    db.commit()
    db.refresh(room)
    log_action(db, "create_room", "room", current_user, {"room_id": room.id, "room_code": room.room_code})
    return {"id": room.id}


@router.put("/rooms/{room_id}")
def update_room(
    room_id: int,
    payload: UpdateRoomRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在"})
    values = payload.dict(exclude_unset=True)
    if "room_code" in values and values["room_code"] != room.room_code:
        if db.query(Room).filter(Room.room_code == values["room_code"], Room.id != room_id).first():
            raise HTTPException(status_code=400, detail={"code": "ROOM_EXISTS", "message": "机房编码已存在"})
    if "qr_content" in values and values["qr_content"] != room.qr_content:
        if db.query(Room).filter(Room.qr_content == values["qr_content"], Room.id != room_id).first():
            raise HTTPException(status_code=400, detail={"code": "ROOM_QR_EXISTS", "message": "二维码值已存在"})
        if db.query(InspectionPoint).filter(InspectionPoint.qr_content == values["qr_content"], InspectionPoint.room_id != room_id).first():
            raise HTTPException(status_code=400, detail={"code": "ROOM_QR_EXISTS", "message": "二维码值已存在"})
    if "nfc_tag" in values and values["nfc_tag"] and values["nfc_tag"] != room.nfc_tag:
        if db.query(Room).filter(Room.nfc_tag == values["nfc_tag"], Room.id != room_id).first():
            raise HTTPException(status_code=400, detail={"code": "ROOM_NFC_EXISTS", "message": "NFC 值已存在"})
        if db.query(InspectionPoint).filter(InspectionPoint.nfc_tag == values["nfc_tag"], InspectionPoint.room_id != room_id).first():
            raise HTTPException(status_code=400, detail={"code": "ROOM_NFC_EXISTS", "message": "NFC 值已存在"})
    if "check_items" in values:
        values["check_items"] = _normalize_check_items(values.get("check_items"))
    _apply_updates(room, values)
    _ensure_room_inspection_point(db, room)
    db.commit()
    db.refresh(room)
    log_action(db, "update_room", "room", current_user, {"room_id": room.id, "room_code": room.room_code})
    return {"id": room.id}


@router.delete("/rooms/{room_id}")
def deactivate_room(
    room_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在"})
    room.is_active = False
    room.updated_at = datetime.utcnow()
    point = db.query(InspectionPoint).filter(InspectionPoint.room_id == room.id, InspectionPoint.point_code == f"ROOM-{room.id}").first()
    if point:
        point.is_active = False
        point.updated_at = datetime.utcnow()
    db.commit()
    log_action(db, "deactivate_room", "room", current_user, {"room_id": room.id, "room_code": room.room_code})
    return {"id": room.id, "is_active": room.is_active}


@router.patch("/rooms/{room_id}/active")
def set_room_active(
    room_id: int,
    is_active: bool,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    room = db.query(Room).filter(Room.id == room_id).first()
    if not room:
        raise HTTPException(status_code=404, detail={"code": "ROOM_NOT_FOUND", "message": "机房不存在"})
    room.is_active = is_active
    room.updated_at = datetime.utcnow()
    _ensure_room_inspection_point(db, room)
    db.commit()
    log_action(db, "set_room_active", "room", current_user, {"room_id": room.id, "room_code": room.room_code, "is_active": is_active})
    return {"id": room.id, "is_active": room.is_active}


@router.get("/inspection-points")
def list_inspection_points(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    items = db.query(InspectionPoint).order_by(InspectionPoint.id.asc()).all()
    room_ids = [item.room_id for item in items if item.room_id is not None]
    system_ids = [item.system_id for item in items if item.system_id is not None]
    rooms = {
        item.id: item
        for item in db.query(Room).filter(Room.id.in_(room_ids)).all()
    } if room_ids else {}
    systems = {
        item.id: item
        for item in db.query(System).filter(System.id.in_(system_ids)).all()
    } if system_ids else {}
    return {
        "items": [
            {
                "id": item.id,
            "room_id": item.room_id,
            "room_name": rooms[item.room_id].room_name if item.room_id in rooms else "",
            "room_weekday_inspection_count": rooms[item.room_id].weekday_inspection_count if item.room_id in rooms else 1,
            "room_holiday_inspection_count": rooms[item.room_id].holiday_inspection_count if item.room_id in rooms else 1,
            "system_id": item.system_id,
                "system_name": systems[item.system_id].name if item.system_id in systems else "",
                "point_code": item.point_code,
                "point_name": item.point_name,
                "point_type": item.point_type,
                "qr_content": item.qr_content,
                "nfc_tag": item.nfc_tag,
                "location_detail": item.location_detail,
                "is_active": item.is_active,
            }
            for item in items
        ]
    }


@router.post("/inspection-points")
def create_inspection_point(
    payload: CreateInspectionPointRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    if db.query(InspectionPoint).filter(InspectionPoint.point_code == payload.point_code).first():
        raise HTTPException(status_code=400, detail={"code": "POINT_EXISTS", "message": "巡检点编码已存在"})
    _validate_inspection_point_refs(db, payload.room_id, payload.system_id)
    point = InspectionPoint(
        room_id=payload.room_id,
        system_id=payload.system_id,
        point_code=payload.point_code,
        point_name=payload.point_name,
        point_type=payload.point_type,
        qr_content=payload.qr_content or payload.point_code,
        nfc_tag=payload.nfc_tag,
        location_detail=payload.location_detail,
        is_active=payload.is_active,
        location=payload.location_detail,
    )
    db.add(point)
    db.commit()
    db.refresh(point)
    log_action(db, "create_inspection_point", "inspection_point", current_user, {"point_id": point.id, "point_code": point.point_code})
    return {"id": point.id}


@router.put("/inspection-points/{point_id}")
def update_inspection_point(
    point_id: int,
    payload: UpdateInspectionPointRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    point = db.query(InspectionPoint).filter(InspectionPoint.id == point_id).first()
    if not point:
        raise HTTPException(status_code=404, detail={"code": "POINT_NOT_FOUND", "message": "巡检点不存在"})
    values = payload.dict(exclude_unset=True)
    if "point_code" in values and values["point_code"] != point.point_code:
        if db.query(InspectionPoint).filter(InspectionPoint.point_code == values["point_code"], InspectionPoint.id != point_id).first():
            raise HTTPException(status_code=400, detail={"code": "POINT_EXISTS", "message": "巡检点编码已存在"})
    if "qr_content" in values and values["qr_content"] != point.qr_content:
        if db.query(InspectionPoint).filter(InspectionPoint.qr_content == values["qr_content"], InspectionPoint.id != point_id).first():
            raise HTTPException(status_code=400, detail={"code": "POINT_QR_EXISTS", "message": "巡检点二维码内容已存在"})
    if "room_id" in values or "system_id" in values:
        _validate_inspection_point_refs(db, values.get("room_id", point.room_id), values.get("system_id", point.system_id))
    if "location_detail" in values:
        values["location"] = values["location_detail"]
    _apply_updates(point, values)
    db.commit()
    db.refresh(point)
    log_action(db, "update_inspection_point", "inspection_point", current_user, {"point_id": point.id, "point_code": point.point_code})
    return {"id": point.id}


@router.delete("/inspection-points/{point_id}")
def deactivate_inspection_point(
    point_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    point = db.query(InspectionPoint).filter(InspectionPoint.id == point_id).first()
    if not point:
        raise HTTPException(status_code=404, detail={"code": "POINT_NOT_FOUND", "message": "巡检点不存在"})
    point.is_active = False
    point.updated_at = datetime.utcnow()
    db.commit()
    log_action(db, "deactivate_inspection_point", "inspection_point", current_user, {"point_id": point.id, "point_code": point.point_code})
    return {"id": point.id, "is_active": point.is_active}


@router.patch("/inspection-points/{point_id}/active")
def set_inspection_point_active(
    point_id: int,
    is_active: bool,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    point = db.query(InspectionPoint).filter(InspectionPoint.id == point_id).first()
    if not point:
        raise HTTPException(status_code=404, detail={"code": "POINT_NOT_FOUND", "message": "巡检点不存在"})
    point.is_active = is_active
    point.updated_at = datetime.utcnow()
    db.commit()
    log_action(
        db,
        "set_inspection_point_active",
        "inspection_point",
        current_user,
        {"point_id": point.id, "point_code": point.point_code, "is_active": is_active},
    )
    return {"id": point.id, "is_active": point.is_active}


@router.get("/inspection-records")
def list_admin_inspection_records(
    page: int = 1,
    size: int = 20,
    system_id: Optional[int] = None,
    room_id: Optional[int] = None,
    point_id: Optional[int] = None,
    result: Optional[str] = None,
    start_at: Optional[datetime] = None,
    end_at: Optional[datetime] = None,
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    page = max(page, 1)
    size = min(max(size, 1), 100)
    q = db.query(InspectionRecord)
    if system_id is not None:
        q = q.filter(InspectionRecord.system_id == system_id)
    if room_id is not None:
        q = q.filter(InspectionRecord.room_id == room_id)
    if point_id is not None:
        q = q.filter(InspectionRecord.point_id == point_id)
    if result:
        q = q.filter(InspectionRecord.result == result)
    if start_at:
        q = q.filter(InspectionRecord.inspected_at >= start_at)
    if end_at:
        q = q.filter(InspectionRecord.inspected_at <= end_at)

    total = q.count()
    records = q.order_by(InspectionRecord.inspected_at.desc()).offset((page - 1) * size).limit(size).all()
    point_ids = [item.point_id for item in records]
    room_ids = [item.room_id for item in records if item.room_id is not None]
    system_ids = [item.system_id for item in records if item.system_id is not None]
    inspector_ids = [item.inspector_id for item in records]

    points = {
        item.id: item
        for item in db.query(InspectionPoint).filter(InspectionPoint.id.in_(point_ids)).all()
    } if point_ids else {}
    rooms = {
        item.id: item
        for item in db.query(Room).filter(Room.id.in_(room_ids)).all()
    } if room_ids else {}
    systems = {
        item.id: item
        for item in db.query(System).filter(System.id.in_(system_ids)).all()
    } if system_ids else {}
    users = {
        item.id: item
        for item in db.query(User).filter(User.id.in_(inspector_ids)).all()
    } if inspector_ids else {}

    items = []
    for record in records:
        point = points.get(record.point_id)
        room = rooms.get(record.room_id)
        system = systems.get(record.system_id)
        inspector = users.get(record.inspector_id)
        note = _parse_inspection_note(record.note)
        items.append({
            "id": record.id,
            "system_id": record.system_id,
            "system_name": system.name if system else "",
            "point_id": record.point_id,
            "point_code": point.point_code if point else "",
            "point_name": point.point_name if point else "",
            "room_id": record.room_id,
            "room_name": room.room_name if room else "",
            "inspector_id": record.inspector_id,
            "inspector_name": inspector.nickname or inspector.username if inspector else "",
            "result": record.result,
            "note": note["note"],
            "check_results": note["check_results"],
            "monitoring_confirmation": note["monitoring_confirmation"],
            "source": record.source,
            "inspected_at": record.inspected_at,
        })

    return {
        "page": page,
        "size": size,
        "total": total,
        "filters": {
            "system_id": system_id,
            "room_id": room_id,
            "point_id": point_id,
            "result": result,
            "start_at": start_at,
            "end_at": end_at,
        },
        "items": items,
    }


@router.post("/assets/batch")
def batch_create_assets(
    payload: BatchCreateAssetsRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("批量创建资产请求: operator=%s item_count=%s", current_user.username, len(payload.items))
    created = []
    skipped = []

    seen_codes: Set[str] = set()
    items_to_create = []
    for item in payload.items:
        code = item.asset_code.strip()
        if code in seen_codes:
            skipped.append({"asset_code": code, "reason": "BATCH_DUPLICATE"})
            continue
        seen_codes.add(code)
        items_to_create.append(item)

    existing_codes = set()
    if items_to_create:
        existing_codes = {
            row[0]
            for row in db.query(Asset.asset_code).filter(Asset.asset_code.in_([it.asset_code for it in items_to_create])).all()
        }

    for item in items_to_create:
        if item.asset_code in existing_codes:
            skipped.append({"asset_code": item.asset_code, "reason": "ASSET_EXISTS"})
            continue
        try:
            _validate_asset_refs(db, item.system_id, item.room_id)
        except HTTPException as exc:
            skipped.append({"asset_code": item.asset_code, "reason": exc.detail.get("code", "INVALID_REF")})
            continue
        asset = Asset(**item.dict())
        db.add(asset)
        db.flush()
        created.append({"id": asset.id, "asset_code": asset.asset_code})

    db.commit()
    logger.info("批量创建资产完成: created=%s skipped=%s", len(created), len(skipped))
    log_action(
        db,
        "batch_create_assets",
        "asset",
        current_user,
        {"created_count": len(created), "skipped_count": len(skipped)},
    )
    return {"created": created, "skipped": skipped}


@router.post("/threat-intel/ip-reputation")
def query_ip_reputation(
    payload: ThreatIntelQueryRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    logger.info("威胁情报批量查询请求: operator=%s count=%s", current_user.username, len(payload.ips))
    try:
        result = batch_query_ip_reputation(payload.ips, lang=payload.lang, realtime_verdict=payload.realtime_verdict)
    except ThreatbookError as exc:
        logger.warning("威胁情报批量查询失败: %s", str(exc))
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    logger.info("威胁情报批量查询完成: total=%s high_risk=%s", result["summary"]["total"], result["summary"]["high_risk"])
    log_action(
        db,
        "query_ip_reputation",
        "threat_intel",
        current_user,
        {
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.post("/threat-intel/ip-reputation/quick")
def query_ip_reputation_quick(
    payload: ThreatIntelQuickInputRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    raw_items = [payload.raw_input]
    try:
        result = batch_query_ip_reputation(raw_items, lang=payload.lang, realtime_verdict=payload.realtime_verdict)
    except ThreatbookError as exc:
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    log_action(
        db,
        "query_ip_reputation_quick",
        "threat_intel",
        current_user,
        {
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.post("/threat-intel/ip-reputation/excel")
async def query_ip_reputation_excel(
    file: UploadFile = File(...),
    lang: str = "zh",
    realtime_verdict: bool = True,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(status_code=400, detail={"code": "INVALID_FILE_TYPE", "message": "仅支持上传 .xlsx 类 Excel 文件"})

    content = await file.read()
    raw_ips = _extract_ips_from_excel(content)
    try:
        result = batch_query_ip_reputation(raw_ips, lang=lang, realtime_verdict=realtime_verdict)
    except ThreatbookError as exc:
        raise HTTPException(status_code=400, detail={"code": "THREATBOOK_QUERY_FAILED", "message": str(exc)})

    result["import"] = {"filename": filename, "source": "excel"}
    log_action(
        db,
        "query_ip_reputation_excel",
        "threat_intel",
        current_user,
        {
            "filename": filename,
            "count": result["summary"]["total"],
            "high_risk": result["summary"]["high_risk"],
            "block_candidates": result["summary"]["block_candidates"],
        },
    )
    return result


@router.get("/emergency-config")
def get_emergency_config(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("super_admin")),
):
    return list_public_config(db)


@router.post("/emergency-config/ssh-hosts")
def save_emergency_ssh_host(
    payload: EmergencySshHostSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_ssh_host(payload, db=db)
    log_action(db, "save_emergency_ssh_host", "emergency_config", current_user, {"host_code": payload.host_code})
    return item


@router.delete("/emergency-config/ssh-hosts/{host_code}")
def delete_emergency_ssh_host(
    host_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = delete_ssh_host(host_code, db=db)
    log_action(db, "delete_emergency_ssh_host", "emergency_config", current_user, {"host_code": host_code})
    return result


@router.post("/emergency-config/server-actions")
def save_emergency_server_action(
    payload: EmergencyServerActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_server_action(payload, db=db)
    log_action(db, "save_emergency_server_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.delete("/emergency-config/server-actions/{action_code}")
def delete_emergency_server_action(
    action_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = delete_server_action(action_code, db=db)
    log_action(db, "delete_emergency_server_action", "emergency_config", current_user, {"action_code": action_code})
    return result


@router.post("/emergency-config/database-actions")
def save_emergency_db_action(
    payload: EmergencyDbActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_db_action(payload, db=db)
    log_action(db, "save_emergency_db_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.delete("/emergency-config/database-actions/{action_code}")
def delete_emergency_db_action(
    action_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = delete_db_action(action_code, db=db)
    log_action(db, "delete_emergency_db_action", "emergency_config", current_user, {"action_code": action_code})
    return result


@router.post("/emergency-config/process-actions")
def save_emergency_process_action(
    payload: EmergencyProcessActionSaveRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    item = upsert_process_action(payload, db=db)
    log_action(db, "save_emergency_process_action", "emergency_config", current_user, {"action_code": payload.action_code})
    return item


@router.delete("/emergency-config/process-actions/{action_code}")
def delete_emergency_process_action(
    action_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = delete_process_action(action_code, db=db)
    log_action(db, "delete_emergency_process_action", "emergency_config", current_user, {"action_code": action_code})
    return result


@router.post("/emergency-config/import-json")
def import_emergency_config_json(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("super_admin")),
):
    result = import_emergency_json_to_db(db)
    log_action(db, "import_emergency_config_json", "emergency_config", current_user, result)
    return result


@router.post("/threat-intel/block-ip")
def block_high_risk_ip(
    payload: ThreatIntelBlockRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    try:
        result = block_ip_with_firewall(ip=payload.ip, reason=payload.reason, dry_run=payload.dry_run, db=db)
    except FirewallConfigError as exc:
        raise HTTPException(status_code=503, detail={"code": "FIREWALL_CONFIG_MISSING", "message": str(exc)}) from exc
    except FirewallClientError as exc:
        raise HTTPException(status_code=502, detail={"code": "FIREWALL_REQUEST_FAILED", "message": str(exc)}) from exc
    except FirewallError as exc:
        raise HTTPException(status_code=400, detail={"code": "FIREWALL_BLOCK_INVALID", "message": str(exc)}) from exc

    result.update(
        {
            "risk_level": payload.risk_level,
            "reason": payload.reason,
            "source": payload.source,
        }
    )
    log_action(db, "block_ip_request", "threat_intel", current_user, result)
    return result


@router.get("/threat-intel/firewall-config")
def get_threat_intel_firewall_config(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles("admin", "super_admin")),
):
    return serialize_firewall_config(db)


@router.put("/threat-intel/firewall-config")
def save_threat_intel_firewall_config(
    payload: FirewallBlockConfigRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles("admin", "super_admin")),
):
    result = upsert_firewall_config(db, payload)
    log_action(
        db,
        "save_firewall_block_config",
        "threat_intel",
        current_user,
        {
            "enabled": result["enabled"],
            "firewall_ip": result["firewall_ip"],
            "address_book_name": result["address_book_name"],
            "has_username": result["has_username"],
            "has_password": result["has_password"],
        },
    )
    return result
